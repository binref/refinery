from __future__ import annotations

import datetime
import io
import struct
import zipfile

from refinery.lib.excel import (
    Cell,
    CellKind,
    ExcelFormat,
    ExcelFormatError,
    SheetKind,
    detect_format,
    open_workbook,
)
from refinery.lib.excel.common import decode_xstring
from refinery.lib.excel.formula import synthesize_formula

from ... import TestBase
from .samples import (
    APACHE_POI_52348,
    BINARY_REFINERY,
    LOWER_CASE_CELLNAMES,
    REVENG1,
    SELF_EVALUATION_REPORT,
    SHARED_STRINGS_ALT_LOCATION,
    XLM_MACRO_FORMULA_XLSM,
    XLM_MACRO_TEXT_XLSM,
)

#: The tail of an XML declaration followed by a document type that declares one entity, which the
#: hardened parser of the reader refuses to expand.
_ENTITY_DECLARATION = b'?><!DOCTYPE root [<!ENTITY e "x">]>'


def _cells(data: bytes) -> dict[tuple[str, int, int], Cell]:
    return {
        (sheet.name, cell.row, cell.col): cell
        for sheet in open_workbook(data).sheets()
        for cell in sheet.cells()
    }


def _entries(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {
            info.filename: archive.read(info)
            for info in archive.infolist()
            if not info.is_dir()
        }


def _package(entries: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w') as archive:
        for name, blob in entries.items():
            archive.writestr(name, blob)
    return buffer.getvalue()


def _normalized(value):
    # openpyxl decodes ST_Xstring escapes in shared strings but leaves them in formula results
    # and inline strings, so both sides are decoded to a fixed point before comparison; integral
    # floats are folded to integers the way the reader casts numbers.
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        while (decoded := decode_xstring(value)) != value:
            value = decoded
    return value


def _replace_part(data: bytes, part: str, replacements: list[tuple[bytes, bytes]]) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(data))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w') as target:
        for info in source.infolist():
            content = source.read(info)
            if info.filename == part:
                for old, new in replacements:
                    content = content.replace(old, new)
            target.writestr(info, content)
    return buffer.getvalue()


def _corrupt_compressed_part(data: bytes, part: str) -> bytes:
    # Flipping a byte inside the deflate stream of an entry makes reading it raise the error
    # of the `zlib` module; the position of the stream is taken from the central directory
    # because the local header of the entry does not carry the compressed size.
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        info = archive.getinfo(part)
    offset = info.header_offset
    name_length, extra_length = struct.unpack_from('<HH', data, offset + 26)
    payload = offset + 30 + name_length + extra_length
    modified = bytearray(data)
    modified[payload + info.compress_size // 2] ^= 0xFF
    return bytes(modified)


class TestWorkbookStructure(TestBase):

    def test_detect_ooxml(self):
        for data in [
            REVENG1,
            LOWER_CASE_CELLNAMES,
            SELF_EVALUATION_REPORT,
            SHARED_STRINGS_ALT_LOCATION,
        ]:
            with self.subTest(data=data[:8]):
                self.assertIs(detect_format(data), ExcelFormat.OOXML)

    def test_rejects_non_workbook(self):
        self.assertIsNone(detect_format(b'not a spreadsheet'))
        self.assertIsNone(detect_format(b'PK\x03\x04' + bytes(64)))

    def test_sheets_of_reveng1(self):
        self.assertEqual(
            [(sheet.name, sheet.kind) for sheet in open_workbook(REVENG1).sheets()],
            [
                ('ZZZfirstsheet', SheetKind.WORKSHEET),
                ('AAA2ndsheet', SheetKind.WORKSHEET),
                ('ControlChars', SheetKind.WORKSHEET),
            ],
        )

    def test_cell_kinds(self):
        cells = _cells(REVENG1)
        zzz = 'ZZZfirstsheet'
        self.assertEqual(cells[(zzz, 1, 1)], Cell(1, 1, CellKind.TEXT, 'description', None))
        self.assertEqual(cells[(zzz, 7, 2)], Cell(7, 2, CellKind.NUMBER, 0, None))
        self.assertEqual(cells[(zzz, 12, 2)], Cell(12, 2, CellKind.NUMBER, 123456, None))
        self.assertEqual(cells[(zzz, 15, 2)], Cell(15, 2, CellKind.BOOLEAN, True, None))
        self.assertEqual(cells[(zzz, 16, 2)], Cell(16, 2, CellKind.BOOLEAN, False, None))
        self.assertEqual(cells[(zzz, 17, 2)], Cell(17, 2, CellKind.ERROR, '#DIV/0!', None))
        self.assertEqual(cells[(zzz, 24, 2)], Cell(24, 2, CellKind.BLANK, None, None))
        self.assertEqual(
            cells[(zzz, 28, 2)],
            Cell(28, 2, CellKind.DATE, datetime.datetime(1999, 12, 31), None),
        )
        self.assertEqual(
            cells[(zzz, 45, 2)],
            Cell(45, 2, CellKind.DATE, datetime.time(2, 10, 54, 545000), '1/11'),
        )

    def test_formula_cells_carry_cached_results(self):
        cells = _cells(REVENG1)
        zzz = 'ZZZfirstsheet'
        # A formula cell with a cached result carries the kind and value of that result together
        # with the uninterpreted formula source.
        self.assertEqual(
            cells[(zzz, 3, 3)],
            Cell(3, 3, CellKind.TEXT, 127 * 'x', 'REPT("x",D3)'),
        )
        self.assertEqual(
            cells[(zzz, 28, 3)],
            Cell(28, 3, CellKind.DATE, datetime.datetime(2000, 1, 1), 'B28+1'),
        )
        # A formula cell without a cached result carries no value.
        self.assertEqual(cells[(zzz, 2, 3)], Cell(2, 3, CellKind.FORMULA, None, '""'))

    def test_inline_strings(self):
        cells = _cells(APACHE_POI_52348)
        # All text of this workbook is stored as inline strings rather than shared strings.
        self.assertEqual(cells[('balance', 2, 3)], Cell(2, 3, CellKind.TEXT, 'Category', None))
        # An inline string cell without an `is` element carries no value.
        self.assertEqual(cells[('balance', 1, 1)], Cell(1, 1, CellKind.BLANK, None, None))

    def test_iso_date_cells(self):
        # A cell of type `d` stores its value as an ISO 8601 date rather than as a serial
        # number whose format marks it as a date.
        data = _replace_part(
            SHARED_STRINGS_ALT_LOCATION,
            'xl/worksheets/sheet1.xml',
            [(
                b'<c r="A2" s="0" t="n"><v>1</v></c>',
                b'<c r="A2" s="0" t="d"><v>2017-12-27T00:00:00</v></c>',
            )],
        )
        self.assertEqual(
            _cells(data)[('Sheet1', 2, 1)],
            Cell(2, 1, CellKind.DATE, datetime.datetime(2017, 12, 27), None),
        )

    def test_workbook_openpyxl_cannot_load(self):
        # openpyxl rejects the styles part of this workbook over an unsupported `builtinId`
        # attribute; the values below are read from the shared string part of the file itself.
        self.assertEqual(
            [(sheet.name, sheet.kind) for sheet in open_workbook(SELF_EVALUATION_REPORT).sheets()],
            [
                ('test1', SheetKind.WORKSHEET),
                ('test2', SheetKind.WORKSHEET),
            ],
        )
        self.assertEqual(_cells(SELF_EVALUATION_REPORT), {
            ('test1', 1, 1): Cell(1, 1, CellKind.TEXT, 'one', None),
            ('test1', 1, 2): Cell(1, 2, CellKind.TEXT, 'two', None),
            ('test2', 1, 1): Cell(1, 1, CellKind.TEXT, 'three', None),
            ('test2', 1, 2): Cell(1, 2, CellKind.TEXT, 'four', None),
        })


class TestNonstandardPackages(TestBase):
    """
    Workbooks written by non-Microsoft software deviate from the conventional part names of the
    package; the reader compensates for each of the deviations below without changing the cells
    it reads.
    """

    def test_uppercase_entry_names(self):
        entries = _entries(SHARED_STRINGS_ALT_LOCATION)
        mangled = {name.upper(): blob for name, blob in entries.items()}
        self.assertEqual(_cells(_package(mangled)), _cells(SHARED_STRINGS_ALT_LOCATION))

    def test_backslash_entry_names(self):
        entries = _entries(SHARED_STRINGS_ALT_LOCATION)
        mangled = {name.replace('/', '\\'): blob for name, blob in entries.items()}
        self.assertEqual(_cells(_package(mangled)), _cells(SHARED_STRINGS_ALT_LOCATION))

    def test_lowercase_content_types(self):
        entries = _entries(SHARED_STRINGS_ALT_LOCATION)
        entries['[content_types].xml'] = entries.pop('[Content_Types].xml')
        self.assertEqual(_cells(_package(entries)), _cells(SHARED_STRINGS_ALT_LOCATION))

    def test_absolute_relationship_targets(self):
        entries = _entries(SHARED_STRINGS_ALT_LOCATION)
        rels = entries['xl/_rels/workbook.xml.rels'].decode()
        entries['xl/_rels/workbook.xml.rels'] = rels.replace(
            'Target="worksheets/sheet1.xml"',
            'Target="/xl/worksheets/sheet1.xml"',
        ).encode()
        self.assertEqual(_cells(_package(entries)), _cells(SHARED_STRINGS_ALT_LOCATION))

    def test_relocated_workbook_part(self):
        entries = _entries(SHARED_STRINGS_ALT_LOCATION)
        entries['xl/main.xml'] = entries.pop('xl/workbook.xml')
        entries['xl/_rels/main.xml.rels'] = entries.pop('xl/_rels/workbook.xml.rels')
        entries['[Content_Types].xml'] = entries['[Content_Types].xml'].replace(
            b'PartName="/xl/workbook.xml"', b'PartName="/xl/main.xml"')
        relocated = _package(entries)
        self.assertIs(detect_format(relocated), ExcelFormat.OOXML)
        self.assertEqual(_cells(relocated), _cells(SHARED_STRINGS_ALT_LOCATION))

    def test_duplicate_part_names_resolve_to_the_last_entry(self):
        # The `zipfile` module resolves a part that is stored twice to its last entry, and so
        # does the package of the reader.
        with zipfile.ZipFile(io.BytesIO(SHARED_STRINGS_ALT_LOCATION)) as source:
            other = source.read('xl/workbook.xml').replace(b'name="Sheet1"', b'name="Last"')
            buffer = io.BytesIO()
            with zipfile.ZipFile(buffer, 'w') as archive:
                for info in source.infolist():
                    archive.writestr(info.filename, source.read(info))
                archive.writestr('xl/workbook.xml', other)
        self.assertEqual(
            [sheet.name for sheet in open_workbook(buffer.getvalue()).sheets()],
            ['Last'],
        )

    def test_corrupt_content_types_fall_back_to_the_default_part(self):
        entries = _entries(SHARED_STRINGS_ALT_LOCATION)
        entries['[Content_Types].xml'] = b'this is not xml'
        self.assertEqual(_cells(_package(entries)), _cells(SHARED_STRINGS_ALT_LOCATION))

    def test_content_types_that_declare_an_entity_fall_back_to_the_default_part(self):
        entries = _entries(SHARED_STRINGS_ALT_LOCATION)
        entries['[Content_Types].xml'] = entries['[Content_Types].xml'].replace(
            b'?>', _ENTITY_DECLARATION)
        self.assertEqual(_cells(_package(entries)), _cells(SHARED_STRINGS_ALT_LOCATION))

    def test_content_types_that_declare_an_entity_locate_no_relocated_workbook_part(self):
        entries = _entries(SHARED_STRINGS_ALT_LOCATION)
        entries['xl/main.xml'] = entries.pop('xl/workbook.xml')
        entries['xl/_rels/main.xml.rels'] = entries.pop('xl/_rels/workbook.xml.rels')
        entries['[Content_Types].xml'] = entries['[Content_Types].xml'].replace(
            b'PartName="/xl/workbook.xml"', b'PartName="/xl/main.xml"').replace(
            b'?>', _ENTITY_DECLARATION)
        self.assertIsNone(detect_format(_package(entries)))


class TestDefectiveReferences(TestBase):
    """
    A row or cell reference that cannot be resolved degrades to the sequential position of
    its element instead of ending the extraction of the sheet.
    """

    def test_unicode_digit_row_reference(self):
        data = _replace_part(
            SHARED_STRINGS_ALT_LOCATION,
            'xl/worksheets/sheet1.xml',
            [(b'<row r="2"', '<row r="²"'.encode())],
        )
        self.assertEqual(_cells(data), _cells(SHARED_STRINGS_ALT_LOCATION))

    def test_zero_row_cell_reference(self):
        data = _replace_part(
            SHARED_STRINGS_ALT_LOCATION,
            'xl/worksheets/sheet1.xml',
            [(b'<c r="A2"', b'<c r="A0"')],
        )
        self.assertEqual(_cells(data), _cells(SHARED_STRINGS_ALT_LOCATION))

    def test_zero_row_reference(self):
        # the cells of the row lose their own references, so only the row places them
        data = _replace_part(
            SHARED_STRINGS_ALT_LOCATION,
            'xl/worksheets/sheet1.xml',
            [
                (b'<row r="2"', b'<row r="0"'),
                (b'<c r="A2" ', b'<c '),
                (b'<c r="B2" ', b'<c '),
            ],
        )
        self.assertEqual(_cells(data), _cells(SHARED_STRINGS_ALT_LOCATION))


class TestOoxmlFormulaCells(TestBase):

    def test_stored_formula_text_decodes_through_the_workbook(self):
        cells = _cells(REVENG1)
        workbook = open_workbook(REVENG1)
        self.assertEqual(
            synthesize_formula(workbook.formula(cells[('ZZZfirstsheet', 3, 3)].formula)),
            'REPT("x",D3)',
        )
        self.assertEqual(
            synthesize_formula(workbook.formula(cells[('ZZZfirstsheet', 45, 2)].formula)),
            '1/11',
        )

    def test_macrosheet_formula_cells_decode(self):
        cells = _cells(XLM_MACRO_FORMULA_XLSM)
        self.assertEqual(cells[('Cdfea', 2, 5)].formula, 'CHAR(113-2)')
        self.assertEqual(cells[('Cdfea', 2, 12)].formula, 'CHAR(71-6)')
        self.assertEqual(cells[('PCWV', 8, 7)].assignment, False)

    def test_an_assignment_formula_carries_the_bx_flag(self):
        # the `bx` attribute of a formula element marks a cell whose text spells an assignment
        # of a value to a name; no sample in the corpus carries one, so the attribute is added
        # to the formula of an authentic cell here
        data = _replace_part(
            XLM_MACRO_FORMULA_XLSM,
            'xl/macrosheets/intlsheet1.xml',
            [(b'<f>FORMULA(', b'<f bx="1">FORMULA(')],
        )
        cells = _cells(data)
        self.assertEqual(cells[('Cdfea', 2, 5)].assignment, False)
        self.assertEqual(cells[('PCWV', 8, 7)].assignment, True)
        self.assertEqual(
            cells[('PCWV', 8, 7)].formula,
            _cells(XLM_MACRO_FORMULA_XLSM)[('PCWV', 8, 7)].formula,
        )

    def test_a_shared_string_formula_without_a_cached_result_is_a_formula_cell(self):
        data = _replace_part(
            XLM_MACRO_TEXT_XLSM,
            'xl/macrosheets/intlsheet1.xml',
            [(
                b'<c r="BG97" s="9" t="str"><f>"..\\iekdhfe.dsk"</f><v>..\\iekdhfe.dsk</v></c>',
                b'<c r="BG97" s="9" t="s"><f bx="1">"..\\iekdhfe.dsk"</f></c>',
            )],
        )
        self.assertEqual(
            _cells(data)[('Doc1', 97, 59)],
            Cell(97, 59, CellKind.FORMULA, None, '"..\\iekdhfe.dsk"', True),
        )


class TestDefectiveWorkbooks(TestBase):
    """
    A defect inside a workbook part raises `ExcelFormatError` from the sheet that contains it,
    after the cells before the defect have been yielded.
    """

    def test_missing_sheet_part(self):
        entries = _entries(SHARED_STRINGS_ALT_LOCATION)
        del entries['xl/worksheets/sheet1.xml']
        sheet = open_workbook(_package(entries)).sheets()[0]
        with self.assertRaises(ExcelFormatError):
            list(sheet.cells())

    def test_cells_before_a_defect_survive(self):
        entries = _entries(SHARED_STRINGS_ALT_LOCATION)
        sheet_xml = entries['xl/worksheets/sheet1.xml']
        cut = sheet_xml.index(b'</row>') + len(b'</row>')
        entries['xl/worksheets/sheet1.xml'] = sheet_xml[:cut]
        sheet = open_workbook(_package(entries)).sheets()[0]
        collected = []
        with self.assertRaises(ExcelFormatError):
            for cell in sheet.cells():
                collected.append((cell.row, cell.col, cell.value))
        self.assertEqual(collected, [
            (1, 1, 'Test'),
            (1, 2, 'Values'),
        ])

    def test_shared_string_index_out_of_range(self):
        entries = _entries(SHARED_STRINGS_ALT_LOCATION)
        entries['xl/worksheets/sheet1.xml'] = entries['xl/worksheets/sheet1.xml'].replace(
            b'<v>0</v>', b'<v>999</v>', 1)
        sheet = open_workbook(_package(entries)).sheets()[0]
        with self.assertRaises(ExcelFormatError):
            list(sheet.cells())

    def test_missing_shared_string_part(self):
        entries = _entries(SHARED_STRINGS_ALT_LOCATION)
        del entries['xl/sharedStrings.xml']
        sheet = open_workbook(_package(entries)).sheets()[0]
        with self.assertRaises(ExcelFormatError):
            list(sheet.cells())

    def test_corrupted_sheet_part(self):
        data = _corrupt_compressed_part(SHARED_STRINGS_ALT_LOCATION, 'xl/worksheets/sheet1.xml')
        sheet = open_workbook(data).sheets()[0]
        with self.assertRaises(ExcelFormatError):
            list(sheet.cells())

    def test_corrupted_workbook_part(self):
        data = _corrupt_compressed_part(SHARED_STRINGS_ALT_LOCATION, 'xl/workbook.xml')
        with self.assertRaises(ExcelFormatError):
            open_workbook(data)

    def test_a_sheet_part_that_declares_an_entity_is_malformed(self):
        data = _replace_part(
            SHARED_STRINGS_ALT_LOCATION,
            'xl/worksheets/sheet1.xml',
            [(b'?>', _ENTITY_DECLARATION)],
        )
        sheet = open_workbook(data).sheets()[0]
        with self.assertRaises(ExcelFormatError):
            list(sheet.cells())

    def test_a_shared_string_part_that_declares_an_entity_is_malformed(self):
        data = _replace_part(
            SHARED_STRINGS_ALT_LOCATION,
            'xl/sharedStrings.xml',
            [(b'?>', _ENTITY_DECLARATION)],
        )
        sheet = open_workbook(data).sheets()[0]
        with self.assertRaises(ExcelFormatError):
            list(sheet.cells())

    def test_a_workbook_part_that_declares_an_entity_is_malformed(self):
        data = _replace_part(
            SHARED_STRINGS_ALT_LOCATION,
            'xl/workbook.xml',
            [(b'?>', _ENTITY_DECLARATION)],
        )
        with self.assertRaises(ExcelFormatError):
            open_workbook(data)

    def test_repeated_reads_of_defective_shared_strings(self):
        # A shared string part that fails to load fails the same way on every read of the
        # sheet, rather than resolving cells against a table that grows with every attempt.
        entries = _entries(SHARED_STRINGS_ALT_LOCATION)
        strings = entries['xl/sharedStrings.xml']
        entry = b'<si><t xml:space="preserve">d</t></si>'
        entries['xl/sharedStrings.xml'] = strings[:strings.index(entry) + len(entry)]
        sheet = open_workbook(_package(entries)).sheets()[0]
        for attempt in range(2):
            with self.subTest(attempt=attempt):
                with self.assertRaises(ExcelFormatError):
                    list(sheet.cells())

    def test_defective_styles_degrade_to_numbers(self):
        entries = _entries(REVENG1)
        entries['xl/styles.xml'] = b'this is not xml'
        cells = _cells(_package(entries))
        # Without the styles part, the date cell of row 28 degrades to its raw serial number.
        self.assertEqual(
            cells[('ZZZfirstsheet', 28, 2)],
            Cell(28, 2, CellKind.NUMBER, 36525, None),
        )
        self.assertEqual(
            cells[('ZZZfirstsheet', 1, 1)],
            Cell(1, 1, CellKind.TEXT, 'description', None),
        )

    def test_styles_that_declare_an_entity_degrade_like_styles_that_are_no_markup(self):
        declared = _entries(REVENG1)
        declared['xl/styles.xml'] = declared['xl/styles.xml'].replace(b'?>', _ENTITY_DECLARATION)
        garbled = _entries(REVENG1)
        garbled['xl/styles.xml'] = b'this is not xml'
        self.assertEqual(_cells(_package(declared)), _cells(_package(garbled)))


class TestAgainstOpenpyxl(TestBase):

    def test_cells_and_sheets_match(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest('the openpyxl oracle is not installed')
        for name, data in [
            ('binary_refinery', BINARY_REFINERY),
            ('reveng1', REVENG1),
            ('lower_case_cellnames', LOWER_CASE_CELLNAMES),
            ('sharedstrings_alt_location', SHARED_STRINGS_ALT_LOCATION),
            ('apachepoi_52348', APACHE_POI_52348),
        ]:
            with self.subTest(sample=name):
                reference = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
                expected = {
                    (sheet.title, cell.row, cell.column): _normalized(cell.value)
                    for sheet in reference.worksheets
                    for row in sheet.iter_rows()
                    for cell in row
                    if cell.value is not None
                }
                workbook = open_workbook(data)
                self.assertEqual(
                    [sheet.name for sheet in workbook.sheets()],
                    reference.sheetnames,
                )
                actual = {
                    key: _normalized(cell.value)
                    for key, cell in _cells(data).items()
                    if cell.kind not in (CellKind.BLANK, CellKind.FORMULA)
                }
                self.assertEqual(actual, expected)
